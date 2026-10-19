"""Deterministic local parsers for already-scanned input artifacts."""

import csv
from dataclasses import dataclass
from pathlib import Path
from typing import Callable, Literal
from zipfile import BadZipFile, ZipFile


_TEXT_CHUNK_SIZE = 64 * 1024
_PARSER_NAME = "builtin-text-csv"
_PARSER_VERSION = "1.0.0"


@dataclass(frozen=True, slots=True)
class XlsxParseLimits:
    """Budgets enforced before and while reading an XLSX archive."""

    max_archive_entries: int = 10_000
    max_uncompressed_bytes: int = 250 * 1024 * 1024
    max_compression_ratio: int = 100
    max_sheets: int = 50
    max_cells: int = 1_000_000
    max_extracted_text_chars: int = 10_000_000


_DEFAULT_XLSX_PARSE_LIMITS = XlsxParseLimits()


@dataclass(frozen=True, slots=True)
class PdfParseLimits:
    """Budgets and quality threshold for PDF text-layer extraction and OCR."""

    max_pages: int = 500
    max_ocr_pages: int = 20
    max_extracted_text_chars: int = 10_000_000
    minimum_ocr_confidence: float = 60.0


_DEFAULT_PDF_PARSE_LIMITS = PdfParseLimits()


@dataclass(frozen=True, slots=True)
class ImageParseLimits:
    max_pixels: int = 25_000_000
    max_extracted_text_chars: int = 2_000_000
    minimum_ocr_confidence: float = 60.0


_DEFAULT_IMAGE_PARSE_LIMITS = ImageParseLimits()


class InputArtifactParseError(RuntimeError):
    """A terminal parsing error with a stable user-safe reason code."""

    def __init__(self, error_code: str) -> None:
        super().__init__(error_code)
        self.error_code = error_code


@dataclass(frozen=True, slots=True)
class ParsedInputArtifact:
    """A parser result before the Worker persists it to object storage."""

    content_kind: Literal["document", "table"]
    manifest: dict[str, object]
    extracted_text_path: Path
    extracted_text_char_count: int


def parse_input_artifact_file(
    *,
    source_path: Path,
    original_filename: str,
    output_directory: Path,
    xlsx_limits: XlsxParseLimits = _DEFAULT_XLSX_PARSE_LIMITS,
    pdf_limits: PdfParseLimits = _DEFAULT_PDF_PARSE_LIMITS,
    pdf_page_renderer: Callable[[Path, int, Path], Path] | None = None,
    image_limits: ImageParseLimits = _DEFAULT_IMAGE_PARSE_LIMITS,
    image_ocr: Callable[[Path], object] | None = None,
) -> ParsedInputArtifact:
    """Parse one supported local file without evaluating formulas or code."""
    if not source_path.is_file():
        raise ValueError("parse source_path must be an existing file")
    suffix = Path(original_filename).suffix.lower()
    output_directory.mkdir(parents=True, exist_ok=True)
    extracted_text_path = output_directory / "extracted.txt"
    if suffix in {".txt", ".md"}:
        char_count = _copy_utf8_text(source_path, extracted_text_path)
        return ParsedInputArtifact(
            content_kind="document",
            manifest={
                "schema_version": "artifact-parse-manifest/v1",
                "parser_name": _PARSER_NAME,
                "parser_version": _PARSER_VERSION,
                "content_kind": "document",
                "content_trust_level": "untrusted",
                "source_format": suffix[1:],
                "warnings": [],
            },
            extracted_text_path=extracted_text_path,
            extracted_text_char_count=char_count,
        )
    if suffix == ".csv":
        char_count = _copy_utf8_text(source_path, extracted_text_path)
        headers, row_count, ragged_row_count, formula_cell_count = _inspect_csv(
            source_path
        )
        return ParsedInputArtifact(
            content_kind="table",
            manifest={
                "schema_version": "artifact-parse-manifest/v1",
                "parser_name": _PARSER_NAME,
                "parser_version": _PARSER_VERSION,
                "content_kind": "table",
                "content_trust_level": "untrusted",
                "source_format": "csv",
                "column_names": headers,
                "data_row_count": row_count,
                "ragged_row_count": ragged_row_count,
                "formula_cell_count": formula_cell_count,
                "formula_execution": "not_performed",
                "warnings": [],
            },
            extracted_text_path=extracted_text_path,
            extracted_text_char_count=char_count,
        )
    if suffix == ".xlsx":
        return _parse_xlsx(
            source_path=source_path,
            output_path=extracted_text_path,
            limits=xlsx_limits,
        )
    if suffix == ".pdf":
        return _parse_pdf(
            source_path=source_path,
            output_path=extracted_text_path,
            limits=pdf_limits,
            page_renderer=pdf_page_renderer,
            image_ocr=image_ocr,
        )
    if suffix in {".png", ".jpg", ".jpeg", ".webp"}:
        return _parse_image(
            source_path=source_path,
            output_path=extracted_text_path,
            declared_suffix=suffix,
            limits=image_limits,
            image_ocr=image_ocr,
        )
    raise InputArtifactParseError("parser_not_available")


def _copy_utf8_text(source_path: Path, destination_path: Path) -> int:
    char_count = 0
    try:
        with source_path.open("r", encoding="utf-8-sig", newline="") as source:
            with destination_path.open("w", encoding="utf-8", newline="") as target:
                while chunk := source.read(_TEXT_CHUNK_SIZE):
                    char_count += len(chunk)
                    target.write(chunk)
    except UnicodeDecodeError as exc:
        destination_path.unlink(missing_ok=True)
        raise InputArtifactParseError("text_decode_failed") from exc
    return char_count


def _inspect_csv(source_path: Path) -> tuple[list[str], int, int, int]:
    try:
        with source_path.open("r", encoding="utf-8-sig", newline="") as source:
            rows = csv.reader(source)
            headers = next(rows, None)
            if headers is None or not any(header.strip() for header in headers):
                raise InputArtifactParseError("csv_header_missing")
            row_count = 0
            ragged_row_count = 0
            formula_cell_count = 0
            for row in rows:
                row_count += 1
                if len(row) != len(headers):
                    ragged_row_count += 1
                formula_cell_count += sum(
                    1 for cell in row if cell.lstrip().startswith("=")
                )
    except UnicodeDecodeError as exc:
        raise InputArtifactParseError("text_decode_failed") from exc
    except csv.Error as exc:
        raise InputArtifactParseError("csv_structure_invalid") from exc
    return headers, row_count, ragged_row_count, formula_cell_count


def _parse_xlsx(
    *,
    source_path: Path,
    output_path: Path,
    limits: XlsxParseLimits,
) -> ParsedInputArtifact:
    """Extract an XLSX workbook without evaluating formulas or external links."""
    _validate_xlsx_limits(limits)
    _inspect_xlsx_archive(source_path, limits)
    try:
        from openpyxl import load_workbook
        from openpyxl.utils.exceptions import InvalidFileException
    except ImportError as exc:
        raise InputArtifactParseError("parser_dependency_unavailable") from exc

    try:
        workbook = load_workbook(
            source_path,
            read_only=True,
            data_only=False,
            keep_links=False,
        )
    except (BadZipFile, InvalidFileException, KeyError, OSError, ValueError) as exc:
        raise InputArtifactParseError("xlsx_workbook_invalid") from exc

    try:
        if len(workbook.worksheets) > limits.max_sheets:
            raise InputArtifactParseError("xlsx_sheet_budget_exceeded")
        sheet_manifests: list[dict[str, object]] = []
        cell_count = 0
        formula_cell_count = 0
        counter = _CharacterBudget(limits.max_extracted_text_chars)
        try:
            with output_path.open("w", encoding="utf-8", newline="") as target:
                writer = csv.writer(_CountingTextWriter(target, counter))
                for worksheet in workbook.worksheets:
                    writer.writerow([f"# sheet: {worksheet.title}"])
                    first_nonempty_row: list[str] | None = None
                    data_row_count = 0
                    for row in worksheet.iter_rows(values_only=True):
                        cell_count += len(row)
                        if cell_count > limits.max_cells:
                            raise InputArtifactParseError("xlsx_cell_budget_exceeded")
                        values = [_xlsx_cell_text(value) for value in row]
                        formula_cell_count += sum(
                            1 for value in values if value.lstrip().startswith("=")
                        )
                        if any(value for value in values):
                            if first_nonempty_row is None:
                                first_nonempty_row = values
                            else:
                                data_row_count += 1
                        writer.writerow(values)
                    sheet_manifests.append(
                        {
                            "name": worksheet.title,
                            "observed_data_row_count": data_row_count,
                            "first_nonempty_row": first_nonempty_row or [],
                        }
                    )
        except Exception:
            output_path.unlink(missing_ok=True)
            raise
    finally:
        workbook.close()

    return ParsedInputArtifact(
        content_kind="table",
        manifest={
            "schema_version": "artifact-parse-manifest/v1",
            "parser_name": "openpyxl-read-only",
            "parser_version": "3.1.5",
            "content_kind": "table",
            "content_trust_level": "untrusted",
            "source_format": "xlsx",
            "sheet_count": len(sheet_manifests),
            "sheets": sheet_manifests,
            "cell_count": cell_count,
            "formula_cell_count": formula_cell_count,
            "formula_execution": "not_performed",
            "warnings": [],
        },
        extracted_text_path=output_path,
        extracted_text_char_count=counter.character_count,
    )


def _validate_xlsx_limits(limits: XlsxParseLimits) -> None:
    if not all(
        isinstance(value, int) and not isinstance(value, bool) and value > 0
        for value in (
            limits.max_archive_entries,
            limits.max_uncompressed_bytes,
            limits.max_compression_ratio,
            limits.max_sheets,
            limits.max_cells,
            limits.max_extracted_text_chars,
        )
    ):
        raise ValueError("xlsx parse limits must be positive integers")


def _inspect_xlsx_archive(source_path: Path, limits: XlsxParseLimits) -> None:
    try:
        with ZipFile(source_path) as archive:
            entries = [entry for entry in archive.infolist() if not entry.is_dir()]
    except BadZipFile as exc:
        raise InputArtifactParseError("xlsx_archive_invalid") from exc
    if len(entries) > limits.max_archive_entries:
        raise InputArtifactParseError("xlsx_archive_entry_budget_exceeded")
    if any(entry.flag_bits & 0x1 for entry in entries):
        raise InputArtifactParseError("xlsx_encrypted_archive")
    uncompressed_bytes = sum(entry.file_size for entry in entries)
    compressed_bytes = sum(entry.compress_size for entry in entries)
    if uncompressed_bytes > limits.max_uncompressed_bytes:
        raise InputArtifactParseError("xlsx_uncompressed_budget_exceeded")
    if compressed_bytes == 0 and uncompressed_bytes > 0:
        raise InputArtifactParseError("xlsx_compression_ratio_exceeded")
    if compressed_bytes and uncompressed_bytes > compressed_bytes * limits.max_compression_ratio:
        raise InputArtifactParseError("xlsx_compression_ratio_exceeded")


def _xlsx_cell_text(value: object) -> str:
    return "" if value is None else str(value)


@dataclass(slots=True)
class _CharacterBudget:
    maximum: int
    character_count: int = 0

    def add(self, value: str) -> None:
        self.character_count += len(value)
        if self.character_count > self.maximum:
            raise InputArtifactParseError("xlsx_text_budget_exceeded")


class _CountingTextWriter:
    def __init__(self, target, counter: _CharacterBudget) -> None:
        self._target = target
        self._counter = counter

    def write(self, value: str) -> int:
        self._counter.add(value)
        return self._target.write(value)


def _parse_pdf(
    *,
    source_path: Path,
    output_path: Path,
    limits: PdfParseLimits,
    page_renderer: Callable[[Path, int, Path], Path] | None,
    image_ocr: Callable[[Path], object] | None,
) -> ParsedInputArtifact:
    """Extract text layers and OCR only pages without one, under bounded budgets."""
    _validate_pdf_limits(limits)
    with source_path.open("rb") as source:
        magic = source.read(5)
    if magic != b"%PDF-":
        raise InputArtifactParseError("pdf_magic_invalid")
    try:
        from pypdf import PdfReader
        from pypdf.errors import PdfReadError
        from app.modules.input_artifacts.processing.pdf_render import PdfPageRenderError
    except ImportError as exc:
        raise InputArtifactParseError("parser_dependency_unavailable") from exc

    try:
        with source_path.open("rb") as source:
            reader = PdfReader(source, strict=True)
            if reader.is_encrypted:
                raise InputArtifactParseError("pdf_encrypted")
            page_count = len(reader.pages)
            if page_count > limits.max_pages:
                raise InputArtifactParseError("pdf_page_budget_exceeded")
            counter = _CharacterBudget(limits.max_extracted_text_chars)
            page_manifests: list[dict[str, int | str | float | None]] = []
            ocr_page_count = 0
            ocr_page_limit_reached = False
            low_confidence_ocr_page_count = 0
            try:
                with output_path.open("w", encoding="utf-8", newline="") as target:
                    for page_number, page in enumerate(reader.pages, start=1):
                        text = page.extract_text() or ""
                        if not text.strip():
                            if page_renderer is not None and ocr_page_count < limits.max_ocr_pages:
                                rendered_path = output_path.parent / f"ocr-page-{page_number}.png"
                                try:
                                    page_renderer(source_path, page_number, rendered_path)
                                    ocr_text, mean_confidence = _extract_ocr_text(
                                        rendered_path, image_ocr
                                    )
                                finally:
                                    rendered_path.unlink(missing_ok=True)
                                ocr_page_count += 1
                                if ocr_text:
                                    ocr_text_quality = (
                                        "low"
                                        if mean_confidence is None
                                        or mean_confidence
                                        < limits.minimum_ocr_confidence
                                        else "usable"
                                    )
                                    if ocr_text_quality == "low":
                                        low_confidence_ocr_page_count += 1
                                    page_content = f"# page: {page_number}\n{ocr_text}\n"
                                    counter.add(page_content)
                                    target.write(page_content)
                                    page_manifests.append(
                                        {
                                            "page_number": page_number,
                                            "text_char_count": len(ocr_text),
                                            "text_extraction_status": "ocr",
                                            "ocr_mean_confidence": mean_confidence,
                                            "ocr_text_quality": ocr_text_quality,
                                        }
                                    )
                                    continue
                                page_manifests.append(
                                    {
                                        "page_number": page_number,
                                        "text_char_count": 0,
                                        "text_extraction_status": "ocr_empty",
                                        "ocr_mean_confidence": mean_confidence,
                                        "ocr_text_quality": "empty",
                                    }
                                )
                                continue
                            if page_renderer is not None:
                                ocr_page_limit_reached = True
                            page_manifests.append(
                                {
                                    "page_number": page_number,
                                    "text_char_count": 0,
                                    "text_extraction_status": "empty",
                                }
                            )
                            continue
                        page_content = f"# page: {page_number}\n{text}\n"
                        counter.add(page_content)
                        target.write(page_content)
                        page_manifests.append(
                            {
                                "page_number": page_number,
                                "text_char_count": len(text),
                                "text_extraction_status": "text",
                            }
                        )
            except Exception:
                output_path.unlink(missing_ok=True)
                raise
    except InputArtifactParseError:
        raise
    except PdfPageRenderError as exc:
        raise InputArtifactParseError(exc.error_code) from exc
    except (PdfReadError, OSError, ValueError) as exc:
        raise InputArtifactParseError("pdf_structure_invalid") from exc

    final_empty_page_count = sum(
        1 for page in page_manifests if page["text_char_count"] == 0
    )
    text_extraction_coverage = (
        "empty"
        if final_empty_page_count == page_count
        else "partial"
        if final_empty_page_count
        else "complete"
    )
    warnings = (
        ["pdf_no_text_layer"]
        if text_extraction_coverage == "empty" and page_renderer is None
        else ["pdf_ocr_no_text"]
        if text_extraction_coverage == "empty"
        else ["pdf_page_without_text"]
        if text_extraction_coverage == "partial" and not ocr_page_limit_reached
        else []
    )
    if ocr_page_limit_reached:
        warnings.append("pdf_ocr_page_budget_exceeded")
    if low_confidence_ocr_page_count:
        warnings.append("pdf_ocr_low_confidence")
    return ParsedInputArtifact(
        content_kind="document",
        manifest={
            "schema_version": "artifact-parse-manifest/v1",
            "parser_name": "pypdf-text",
            "parser_version": "6.18.1",
            "content_kind": "document",
            "content_trust_level": "untrusted",
            "source_format": "pdf",
            "page_count": page_count,
            "ocr_page_count": ocr_page_count,
            "pages": page_manifests,
            "text_extraction_coverage": text_extraction_coverage,
            "content_verification": "not_performed",
            "warnings": warnings,
        },
        extracted_text_path=output_path,
        extracted_text_char_count=counter.character_count,
    )


def _validate_pdf_limits(limits: PdfParseLimits) -> None:
    if not all(
        isinstance(value, int) and not isinstance(value, bool) and value > 0
        for value in (
            limits.max_pages,
            limits.max_ocr_pages,
            limits.max_extracted_text_chars,
        )
    ):
        raise ValueError("pdf parse limits must be positive integers")
    if (
        not isinstance(limits.minimum_ocr_confidence, (int, float))
        or isinstance(limits.minimum_ocr_confidence, bool)
        or not 0 <= limits.minimum_ocr_confidence <= 100
    ):
        raise ValueError("pdf minimum OCR confidence must be between 0 and 100")


def _parse_image(
    *,
    source_path: Path,
    output_path: Path,
    declared_suffix: str,
    limits: ImageParseLimits,
    image_ocr: Callable[[Path], object] | None,
) -> ParsedInputArtifact:
    _validate_image_limits(limits)
    try:
        from PIL import Image, UnidentifiedImageError
    except ImportError as exc:
        raise InputArtifactParseError("parser_dependency_unavailable") from exc
    expected_format = {
        ".png": "PNG",
        ".jpg": "JPEG",
        ".jpeg": "JPEG",
        ".webp": "WEBP",
    }[declared_suffix]
    try:
        import warnings

        with warnings.catch_warnings():
            warnings.simplefilter("error", Image.DecompressionBombWarning)
            with Image.open(source_path, formats=("PNG", "JPEG", "WEBP")) as image:
                image.verify()
            with Image.open(source_path, formats=("PNG", "JPEG", "WEBP")) as image:
                if image.format != expected_format:
                    raise InputArtifactParseError("image_format_mismatch")
                width, height = image.size
                if width * height > limits.max_pixels:
                    raise InputArtifactParseError("image_pixel_budget_exceeded")
                normalized_path = output_path.parent / "ocr-input.png"
                _save_ocr_png(image, normalized_path, Image)
    except InputArtifactParseError:
        raise
    except (Image.DecompressionBombError, Image.DecompressionBombWarning):
        raise InputArtifactParseError("image_pixel_budget_exceeded") from None
    except (UnidentifiedImageError, OSError, ValueError):
        raise InputArtifactParseError("image_structure_invalid") from None

    try:
        text, mean_confidence = _extract_ocr_text(normalized_path, image_ocr)
    finally:
        normalized_path.unlink(missing_ok=True)

    normalized_text = text.strip()
    if len(normalized_text) > limits.max_extracted_text_chars:
        raise InputArtifactParseError("image_ocr_text_budget_exceeded")
    quality = (
        "empty"
        if not normalized_text
        else "low"
        if mean_confidence is None or mean_confidence < limits.minimum_ocr_confidence
        else "usable"
    )
    warnings = (
        ["image_ocr_no_text"]
        if quality == "empty"
        else ["image_ocr_low_confidence"]
        if quality == "low"
        else []
    )
    try:
        output_path.write_text(normalized_text, encoding="utf-8", newline="")
    except OSError as exc:
        output_path.unlink(missing_ok=True)
        raise InputArtifactParseError("parse_output_write_failed") from exc
    return ParsedInputArtifact(
        content_kind="document",
        manifest={
            "schema_version": "artifact-parse-manifest/v1",
            "parser_name": "tesseract-ocr",
            "content_kind": "document",
            "content_trust_level": "untrusted",
            "source_format": expected_format.lower(),
            "width": width,
            "height": height,
            "ocr_text_quality": quality,
            "ocr_mean_confidence": mean_confidence,
            "warnings": warnings,
        },
        extracted_text_path=output_path,
        extracted_text_char_count=len(normalized_text),
    )


def _extract_ocr_text(
    image_path: Path,
    image_ocr: Callable[[Path], object] | None,
) -> tuple[str, float | None]:
    try:
        from app.modules.input_artifacts.processing.image_ocr import ImageOcrError, TesseractImageOcr

        result = (
            image_ocr(image_path)
            if image_ocr is not None
            else TesseractImageOcr().extract(image_path)
        )
        text = getattr(result, "text")
        mean_confidence = getattr(result, "mean_confidence")
        if not isinstance(text, str):
            raise TypeError("OCR text must be a string")
        if mean_confidence is not None and not isinstance(mean_confidence, float):
            raise TypeError("OCR confidence must be a float or null")
        return text.strip(), mean_confidence
    except ImageOcrError as exc:
        raise InputArtifactParseError(exc.error_code) from exc
    except (OSError, TypeError, ValueError) as exc:
        raise InputArtifactParseError("ocr_engine_failed") from exc


def _save_ocr_png(image, destination_path: Path, image_module) -> None:
    if "A" in image.getbands():
        background = image_module.new("RGBA", image.size, "white")
        background.alpha_composite(image.convert("RGBA"))
        normalized = background.convert("RGB")
    else:
        normalized = image.convert("RGB")
    normalized.save(destination_path, format="PNG")


def _validate_image_limits(limits: ImageParseLimits) -> None:
    if (
        not all(
            isinstance(value, int) and not isinstance(value, bool) and value > 0
            for value in (limits.max_pixels, limits.max_extracted_text_chars)
        )
        or not 0 <= limits.minimum_ocr_confidence <= 100
    ):
        raise ValueError("image parse limits are invalid")
